"""Build editable Word and a readable PDF from the reviewed viva guide.

Optional document-only dependencies: python-docx, reportlab, pypdf.
No change to application dependency locks is required.
"""
from pathlib import Path
from html import escape
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/'POLARIS_Viva_Study_Guide.md'
text=SOURCE.read_text()
lines=text.splitlines()
blocks=[];i=0
while i<len(lines):
    s=lines[i].strip()
    if not s:i+=1;continue
    if s.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in row):rows.append(row)
            i+=1
        blocks.append(('table',rows));continue
    if s.startswith('# '):blocks.append(('h1',s[2:]))
    elif s.startswith('## '):blocks.append(('h2',s[3:]))
    elif s.startswith('- '):blocks.append(('bullet',s[2:]))
    elif re.match(r'^\d+\. ',s):blocks.append(('number',s))
    elif re.match(r'^\*\*Q\d+',s):blocks.append(('question',s))
    else:blocks.append(('p',s))
    i+=1

headings=[v for k,v in blocks if k=='h1'][1:]
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.12
for name,size,col in [('Title',30,'18344B'),('Heading 1',21,'18344B'),('Heading 2',14,'087E8B')]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor.from_string(col)
    style.paragraph_format.keep_with_next=True
    style.paragraph_format.space_before=Pt(12)
    style.paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
head=sec.header.paragraphs[0]
head.text='POLARIS  |  GROUP 18  |  PRESENTATION & VIVA PREPARATION'
head.style='Caption';head.runs[0].font.size=Pt(8);head.runs[0].font.color.rgb=RGBColor.from_string('667789')
foot=sec.footer.paragraphs[0];foot.alignment=2
foot.add_run('Study guide • 14 October 2026  |  Page ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='POLARIS — Presentation and Viva Study Guide'
doc.core_properties.subject='Submitted-slide explanations, components, literature review, and 96 viva questions'
doc.core_properties.author='POLARIS Group 18'

def inline_doc(par,s):
    for j,piece in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
        run=par.add_run(piece);run.bold=(j%2==1)
    par.paragraph_format.widow_control=True

def add_bookmark(par,name,number):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(number));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(number))
    par._p.insert(0,a);par._p.append(b)

def contents_doc():
    doc.add_page_break()
    p=doc.add_paragraph('Study map','Heading 2')
    p=doc.add_paragraph('Present only the submitted slides. Sections 5 and 13 are optional technical backup. Start with sections 1, 2, 7, 11, 14 and the CORE questions if time is short.')
    for n,h in enumerate(headings,1):
        pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(8)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'section_{n}')
        r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'087E8B');rp.append(c);r.append(rp)
        t=OxmlElement('w:t');t.text=h;r.append(t);link.append(r);pp._p.append(link)
    doc.add_paragraph('Use Word’s Navigation Pane to jump between headings. All technical specifications are proposed choices unless explicitly identified as completed software evidence.')

first=True;section=0;added_contents=False
for kind,val in blocks:
    if kind=='h1':
        if first:
            doc.add_paragraph(val,'Title');first=False
        else:
            if not added_contents:contents_doc();added_contents=True
            section+=1;pp=doc.add_paragraph(val,'Heading 1');add_bookmark(pp,f'section_{section}',section)
    elif kind=='h2':doc.add_paragraph(val,'Heading 2')
    elif kind=='table':
        t=doc.add_table(rows=1,cols=len(val[0]));t.style='Light Shading Accent 1'
        t.autofit=False
        for j,v in enumerate(val[0]):inline_doc(t.rows[0].cells[j].paragraphs[0],v)
        trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for row in val[1:]:
            cells=t.add_row().cells
            for j,v in enumerate(row):inline_doc(cells[j].paragraphs[0],v)
        for rid,row in enumerate(t.rows):
            props=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');props.append(ns)
            for cell in row.cells:
                cell.width=Inches(6.87/len(val[0]))
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.line_spacing=1.03
                    for run in pp.runs:run.font.size=Pt(9.5);run.bold=(rid==0) or run.bold
        doc.add_paragraph()
    else:
        pp=doc.add_paragraph(style='List Bullet' if kind=='bullet' else None)
        inline_doc(pp,val)
        if kind=='question':
            pp.paragraph_format.keep_with_next=True
            pp.paragraph_format.space_before=Pt(10)
            for r in pp.runs:r.font.color.rgb=RGBColor.from_string('18344B')
doc.save(ROOT/'POLARIS_Viva_Study_Guide.docx')

# PDF reading copy; same reviewed source text, independently paginated.
fontdir=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('DV',str(fontdir/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DV-Bold',str(fontdir/'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFontFamily('DV',normal='DV',bold='DV-Bold',italic='DV',boldItalic='DV-Bold')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyDV',fontName='DV',fontSize=10.2,leading=14,spaceAfter=7,allowWidows=0,allowOrphans=0))
styles.add(ParagraphStyle(name='TitleDV',fontName='DV-Bold',fontSize=28,leading=35,textColor=colors.HexColor('#18344B'),spaceAfter=20))
styles.add(ParagraphStyle(name='H1DV',fontName='DV-Bold',fontSize=20,leading=26,textColor=colors.HexColor('#18344B'),spaceAfter=15,keepWithNext=True))
styles.add(ParagraphStyle(name='H2DV',fontName='DV-Bold',fontSize=13.4,leading=18,textColor=colors.HexColor('#087E8B'),spaceBefore=11,spaceAfter=8,keepWithNext=True))
styles.add(ParagraphStyle(name='QDV',parent=styles['BodyDV'],fontName='DV-Bold',spaceBefore=10,keepWithNext=True,textColor=colors.HexColor('#18344B')))
styles.add(ParagraphStyle(name='BulletDV',parent=styles['BodyDV'],leftIndent=13,firstLineIndent=-10))
styles.add(ParagraphStyle(name='CellDV',fontName='DV',fontSize=8.8,leading=12.2,spaceAfter=2))
styles.add(ParagraphStyle(name='CellHeadDV',parent=styles['CellDV'],fontName='DV-Bold',textColor=colors.white))

def markup(s):
    return ''.join(('<b>'+escape(v)+'</b>') if i%2 else escape(v) for i,v in enumerate(re.split(r'\*\*(.*?)\*\*',s)))
flow=[];first=True;added_contents=False
for kind,val in blocks:
    if kind=='h1':
        if first:flow.append(Paragraph(markup(val),styles['TitleDV']));first=False
        else:
            if not added_contents:
                flow.extend([PageBreak(),Paragraph('Study map',styles['H1DV']),Paragraph('Present only the submitted slides. Sections 5 and 13 are optional backup. For a short revision, prioritize sections 1, 2, 7, 11, 14 and the CORE viva questions.',styles['BodyDV'])])
                for h in headings:flow.append(Paragraph(markup(h),styles['BodyDV']))
                added_contents=True
            flow.extend([PageBreak(),Paragraph(markup(val),styles['H1DV'])])
    elif kind=='h2':flow.append(Paragraph(markup(val),styles['H2DV']))
    elif kind=='table':
        data=[[Paragraph(markup(v),styles['CellHeadDV'] if i==0 else styles['CellDV']) for v in row] for i,row in enumerate(val)]
        n=len(val[0]);width=A4[0]-100
        widths=([width*.23,width*.34,width*.43] if n==3 else [width/n]*n)
        table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#18344B')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F0F5F8'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#18344B'))]))
        flow.extend([table,Spacer(1,10)])
    else:
        sty='QDV' if kind=='question' else 'BulletDV' if kind=='bullet' else 'BodyDV'
        flow.append(Paragraph(('• ' if kind=='bullet' else '')+markup(val),styles[sty]))

class GuidePDF(SimpleDocTemplate):
    def afterFlowable(self,f):
        if isinstance(f,Paragraph) and f.style.name=='H1DV':
            title=f.getPlainText();key='h'+str(self.page)
            self.canv.bookmarkPage(key);self.canv.addOutlineEntry(title,key,0,False)
def footer(c,d):
    c.saveState();c.setFont('DV',8);c.setFillColor(colors.HexColor('#667789'))
    c.drawString(50,A4[1]-28,'POLARIS • GROUP 18 • PRESENTATION & VIVA PREPARATION')
    c.drawString(50,27,'Submitted slides only • Hardware and AI remain proposed')
    c.drawRightString(A4[0]-50,27,str(d.page));c.restoreState()
pdf=GuidePDF(str(ROOT/'POLARIS_Viva_Study_Guide.pdf'),pagesize=A4,leftMargin=50,rightMargin=50,topMargin=49,bottomMargin=48,title='POLARIS Presentation and Viva Study Guide',author='POLARIS Group 18')
pdf.build(flow,onFirstPage=footer,onLaterPages=footer)
print('Created Word and PDF study guides.')
